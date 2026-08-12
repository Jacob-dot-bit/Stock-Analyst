"""Parser for XTB xStation account exports (XLSX / CSV).

Background: XTB shut down its API on 14 March 2025. Exporting a file is now the only
reliable and terms-compliant way to retrieve your positions. In xStation the path is
*Account history* → *Export*.

Real structure observed on 2026 production exports — one sheet per section::

    Sheet "Open Positions"
        Account number | 1234567
        Open Positions
        Data as of report generated | ...
        Product | Metric | Amount | Currency        <- small summary table
        My Trades | Value | 9344.16 | EUR
        ...
        Product | Instrument/Position | Ticker | Category | Type | Volume | ...
        My Trades | ASML       | ASML.NL | STOCK |      | 1.0 | ...   <- aggregate row
        My Trades | 1636247573 | ASML.NL |       | BUY  | 1.0 | ...   <- lot

Four traps this module handles explicitly, each found the hard way on real files:

1. **Wrong XLSX dimension.** These workbooks declare ``A1:A1``. In ``read_only`` mode
   openpyxl trusts that metadata and returns a single cell, so the file looks empty.
   The workbook is therefore loaded in normal mode.
2. **``Ticker`` is the symbol, ``Instrument`` is the company name.** Confusing the two
   turns "Canadian Pacific" into a ticker.
3. **Open positions come in two levels**: one aggregate row per holding, followed by
   one row per lot. Summing both counts every holding twice.
4. **``Position ID`` is not unique** on closed positions: a holding closed in several
   parts produces several rows sharing one id.

The parser stays tolerant: columns are matched through normalised aliases (English and
French), tables are classified by *column signature* rather than by heading, and
anything not understood is reported in ``warnings`` with the source row preserved in
``raw``. If an export fails to parse, ``COLUMN_ALIASES`` is what needs extending.

User-facing text never appears here. Diagnostics are emitted as ``Message`` codes and
rendered by the client in the user's language.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.messages import Message, MessageCode, SectionKind, SectionSummary
from app.models import TxType

# --- Canonical columns and their aliases ------------------------------------
# Compared after normalisation: lowercase, no accents, no punctuation.
# "%" becomes "pct" so that "Net Profit" and "Net Profit %" stay distinct.

COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    # Instrument identity.
    # NOTE: "Ticker" holds the symbol, "Instrument" holds the company name.
    "symbol": ("ticker", "symbol", "symbole"),
    "name": ("instrument", "instrumentposition", "nom", "libelle", "designation"),
    "category": ("category", "categorie", "classe"),
    "account": ("product", "produit", "compte", "account"),
    # Identifiers
    "position_id": ("positionid", "position", "ticket", "idposition"),
    "cash_id": ("id", "idoperation", "operationid"),
    # Direction and size
    "type": ("type", "direction", "side", "sens", "typedoperation"),
    "volume": ("volume", "quantite", "quantity", "qty", "lots", "nombredeparts"),
    # Prices and dates
    "open_time": ("opentime", "opentimeutc", "heuredouverture", "dateouverture", "ouverture"),
    "open_price": ("openprice", "prixdouverture", "coursdouverture", "prixouverture"),
    "close_time": ("closetime", "closetimeutc", "heuredefermeture", "datefermeture", "fermeture"),
    "close_price": ("closeprice", "prixdefermeture", "coursdecloture", "prixfermeture"),
    "market_price": ("currentprice", "marketprice", "prixdumarche", "coursactuel"),
    # Amounts
    "market_value": ("value", "valeur", "valeurdemarche", "marketvalue"),
    "purchase_value": ("purchasevalue", "valeurdachat", "montantachat"),
    "sale_value": ("salevalue", "valeurdevente", "montantvente"),
    "gross_pl": ("grossprofit", "grosspl", "profitbrut", "plbrut", "resultatbrut"),
    "net_pl": ("netprofit", "profitloss", "netpl", "profitnet", "plnet", "resultatnet", "profit"),
    "net_pl_pct": ("netprofitpct", "profitpct", "plpct", "rendement"),
    "amount": ("amount", "montant"),
    "commission": ("commission", "opencommission", "commissions", "frais"),
    "swap": ("swap", "swaps", "pointsdeswap"),
    "margin": ("margin", "marge"),
    "rollover": ("rollover",),
    # Broker-applied FX rates: needed to reconcile a price quoted in the instrument's
    # currency against an amount expressed in the account currency.
    "open_fx_rate": ("openconversionrate", "tauxdechangeouverture"),
    "close_fx_rate": ("closeconversionrate", "tauxdechangefermeture"),
    "close_origin": ("closeorigin", "originefermeture"),
    # Misc
    "sl": ("sl", "stoploss"),
    "tp": ("tp", "takeprofit"),
    "time": ("time", "timeutc", "heure", "date", "dateheure", "datetime"),
    "comment": ("comment", "commentaire", "description", "remarque"),
    "currency": ("currency", "devise"),
    "metric": ("metric", "metrique", "indicateur"),
}

#: Aliases that are ambiguous between canonical columns; re-arbitrated per table.
_AMBIGUOUS = {"id", "position"}


def _normalize(value: Any) -> str:
    """Lowercase, accent-free, alphanumeric only. "%" is preserved as "pct"."""
    if value is None:
        return ""
    text = str(value).strip().lower().replace("%", "pct")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]", "", text)


_ALIAS_LOOKUP: dict[str, str] = {}
for _canonical, _aliases in COLUMN_ALIASES.items():
    for _alias in _aliases:
        _ALIAS_LOOKUP.setdefault(_normalize(_alias), _canonical)


def _match_column(header: Any) -> str | None:
    return _ALIAS_LOOKUP.get(_normalize(header))


# --- Value conversion --------------------------------------------------------

#: Isolates the first number in a string, thousands separators included.
#: Extracting beats character filtering: otherwise "1 234,56 EUR" leaves the "E"
#: of the currency behind and the conversion fails.
_NUMBER_TOKEN = re.compile(r"[-+]?\d[\d\s.,]*")


def parse_number(value: Any) -> float | None:
    """Parse a number tolerating both English and French conventions.

    Handles "1 234,56", "1,234.56", non-breaking spaces, trailing currency codes and
    accounting parentheses. Returns ``None`` rather than raising: one unreadable cell
    must not fail the whole import.

    Scientific notation is deliberately unsupported — absent from broker exports, it
    would make trailing-currency detection ambiguous.
    """
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).replace("\xa0", " ").replace(" ", " ").strip()
    if not text:
        return None

    # Accounting notation: "(1 234,56)" means -1234.56
    negated = text.startswith("(") and text.endswith(")")
    if negated:
        text = text[1:-1].strip()

    match = _NUMBER_TOKEN.search(text)
    if match is None:
        return None

    text = match.group(0).replace(" ", "").rstrip(".,")
    if not text or text in {"-", "+"}:
        return None

    has_comma, has_dot = "," in text, "." in text
    if has_comma and has_dot:
        # The decimal separator is whichever appears last.
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif has_comma:
        # Comma alone: decimal ("12,5") or thousands separator ("1,234")?
        if re.fullmatch(r"-?\d{1,3}(,\d{3})+", text):
            text = text.replace(",", "")
        else:
            text = text.replace(",", ".")

    try:
        number = float(text)
    except ValueError:
        return None

    return -number if negated else number


_DATE_FORMATS = (
    "%Y-%m-%d %H:%M:%S.%f",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d",
    "%d.%m.%Y %H:%M:%S",
    "%d.%m.%Y %H:%M",
    "%d.%m.%Y",
    "%d/%m/%Y %H:%M:%S",
    "%d/%m/%Y %H:%M",
    "%d/%m/%Y",
    "%m/%d/%Y %H:%M:%S",
    "%m/%d/%Y",
)


def parse_datetime(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value

    text = str(value).strip()
    if not text:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


# --- Cash operation classification -------------------------------------------
# Order matters: "Free funds interest tax" must land on TAX, not INTEREST.

_CASH_TYPE_RULES: tuple[tuple[str, tuple[str, ...]], ...] = (
    (TxType.TAX, ("withholdingtax", "tax", "impot", "taxe", "prelevement", "stampduty", "ifft", "iftt")),
    (TxType.DIVIDEND, ("dividend", "dividende")),
    (TxType.BUY, ("purchase", "achat", "buy")),
    (TxType.SELL, ("sell", "sale", "vente")),
    (TxType.CLOSED_TRADE, ("closetrade", "closedtrade")),
    (TxType.FEE, ("fee", "commission", "frais", "swap", "rollover")),
    (TxType.INTEREST, ("interest", "interet")),
    (TxType.DEPOSIT, ("deposit", "depot", "versement")),
    (TxType.WITHDRAWAL, ("withdrawal", "retrait")),
)

#: Total / subtotal row labels to discard — these are not operations.
_TOTAL_LABELS = {"total", "totaux", "sum", "somme", "profitloss", "profitperte"}


def classify_cash_type(label: Any) -> str:
    normalized = _normalize(label)
    if not normalized:
        return TxType.OTHER
    for tx_type, keywords in _CASH_TYPE_RULES:
        if any(keyword in normalized for keyword in keywords):
            return tx_type
    return TxType.OTHER


def is_total_row(label: Any) -> bool:
    return _normalize(label) in _TOTAL_LABELS


# --- Result structures -------------------------------------------------------


@dataclass
class ParsedTable:
    kind: str
    sheet: str
    columns: dict[int, str]
    unmapped_columns: list[str]
    rows: list[dict[str, Any]]


@dataclass
class ParsedExport:
    open_positions: list[dict[str, Any]] = field(default_factory=list)
    closed_positions: list[dict[str, Any]] = field(default_factory=list)
    cash_operations: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[Message] = field(default_factory=list)
    sections: list[SectionSummary] = field(default_factory=list)
    #: Accounts (the "Product" column) present in the open positions. Used to replace
    #: only the snapshot of the accounts actually re-imported.
    accounts: list[str] = field(default_factory=list)

    @property
    def is_empty(self) -> bool:
        return not (self.open_positions or self.closed_positions or self.cash_operations)


# --- Table classification ----------------------------------------------------


def _classify_table(columns: set[str], sheet_name: str) -> str:
    """Identify a table from its columns.

    The column signature wins over the sheet name: it does not depend on the
    language of the export.
    """
    sheet = _normalize(sheet_name)

    if "symbol" in columns and ("open_price" in columns or "volume" in columns):
        closed_by_columns = "close_price" in columns or "close_time" in columns
        closed_by_sheet = any(word in sheet for word in ("closed", "ferme", "cloture"))
        return SectionKind.CLOSED_POSITIONS if (closed_by_columns or closed_by_sheet) else SectionKind.OPEN_POSITIONS

    if "amount" in columns and "type" in columns:
        return SectionKind.CASH_OPERATIONS

    return SectionKind.UNKNOWN


def _resolve_ambiguous_columns(columns: dict[int, str], raw_headers: dict[int, str]) -> dict[int, str]:
    """Re-arbitrate ambiguous aliases ("ID", "Position") from the other columns."""
    values = set(columns.values())
    is_position_table = "open_price" in values or "close_price" in values

    resolved = dict(columns)
    for index in columns:
        if _normalize(raw_headers.get(index)) in _AMBIGUOUS:
            resolved[index] = "position_id" if is_position_table else "cash_id"
    return resolved


# --- File reading ------------------------------------------------------------


def _rows_from_xlsx(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    # read_only=False on purpose: XTB exports declare a wrong "A1:A1" dimension,
    # which read-only mode trusts, making the workbook look empty.
    workbook = load_workbook(io.BytesIO(content), data_only=True)
    sheets = [
        (worksheet.title, [list(row) for row in worksheet.iter_rows(values_only=True)])
        for worksheet in workbook.worksheets
    ]
    workbook.close()
    return sheets


def _rows_from_csv(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    text = content.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [("csv", [list(row) for row in reader])]


def _is_blank(row: list[Any]) -> bool:
    return all(cell is None or str(cell).strip() == "" for cell in row)


def _extract_tables(sheet_name: str, rows: list[list[Any]]) -> list[ParsedTable]:
    """Scan a sheet and pull out every recognisable table."""
    tables: list[ParsedTable] = []
    index = 0

    while index < len(rows):
        row = rows[index]
        if _is_blank(row):
            index += 1
            continue

        matched_columns = {i: name for i, cell in enumerate(row) if (name := _match_column(cell))}

        # A header row must match enough columns AND yield a usable table. Otherwise
        # it is treated as a preamble line ("Account number | 1234567") or a summary
        # table, and skipped.
        raw_headers = {i: str(cell) for i, cell in enumerate(row) if cell is not None}
        columns = _resolve_ambiguous_columns(matched_columns, raw_headers)
        kind = _classify_table(set(columns.values()), sheet_name)

        if len(matched_columns) < 4 or kind == SectionKind.UNKNOWN:
            index += 1
            continue

        unmapped = [
            str(cell).strip()
            for i, cell in enumerate(row)
            if cell is not None and str(cell).strip() and i not in matched_columns
        ]

        data_rows: list[dict[str, Any]] = []
        cursor = index + 1
        while cursor < len(rows):
            data_row = rows[cursor]
            if _is_blank(data_row):
                break
            # A new header row ends the current table.
            if sum(1 for cell in data_row if _match_column(cell)) >= 4:
                break

            record: dict[str, Any] = {}
            raw: dict[str, Any] = {}
            for col_index, cell in enumerate(data_row):
                header = raw_headers.get(col_index)
                if header is not None and cell is not None:
                    raw[header] = cell if isinstance(cell, (int, float, str)) else str(cell)
                canonical = columns.get(col_index)
                if canonical:
                    record[canonical] = cell
            record["_raw"] = raw
            data_rows.append(record)
            cursor += 1

        tables.append(
            ParsedTable(
                kind=kind,
                sheet=sheet_name,
                columns=columns,
                unmapped_columns=unmapped,
                rows=data_rows,
            )
        )
        index = cursor

    return tables


# --- Entry point -------------------------------------------------------------


def parse_xtb_export(content: bytes, filename: str) -> ParsedExport:
    """Parse an xStation export and return normalised records."""
    result = ParsedExport()
    suffix = Path(filename).suffix.lower()

    try:
        if suffix in {".xlsx", ".xlsm"}:
            sheets = _rows_from_xlsx(content)
        elif suffix in {".csv", ".txt"}:
            sheets = _rows_from_csv(content)
        else:
            result.warnings.append(
                Message(MessageCode.UNSUPPORTED_FILE_TYPE, {"extension": suffix or "?"})
            )
            return result
    except Exception as exc:  # noqa: BLE001 - surfaced to the user, never swallowed
        result.warnings.append(Message(MessageCode.FILE_UNREADABLE, {"error": str(exc)}))
        return result

    tables: list[ParsedTable] = []
    for sheet_name, rows in sheets:
        tables.extend(_extract_tables(sheet_name, rows))

    if not tables:
        result.warnings.append(Message(MessageCode.NO_TABLE_RECOGNISED))
        return result

    for table in tables:
        if table.unmapped_columns:
            result.warnings.append(
                Message(
                    MessageCode.UNMAPPED_COLUMNS,
                    {"sheet": table.sheet, "columns": table.unmapped_columns},
                )
            )

        if table.kind == SectionKind.OPEN_POSITIONS:
            positions = _build_open_positions(table, result)
            result.open_positions.extend(positions)
            count = len(positions)
        elif table.kind == SectionKind.CLOSED_POSITIONS:
            closed = _build_closed_positions(table, result)
            result.closed_positions.extend(closed)
            count = len(closed)
        else:
            operations = _build_cash_operations(table)
            result.cash_operations.extend(operations)
            count = len(operations)

        result.sections.append(
            SectionSummary(
                sheet=table.sheet,
                kind=table.kind,
                count=count,
                source_rows=len(table.rows),
            )
        )

    result.accounts = sorted({p["account"] for p in result.open_positions if p.get("account")})

    if result.is_empty:
        result.warnings.append(Message(MessageCode.NOTHING_IMPORTABLE))

    return result


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _clean_id(value: Any) -> str | None:
    """Normalise a numeric identifier.

    openpyxl returns whole numbers from a workbook as floats: without this, the id
    1677685567 would become the string "1677685567.0".
    """
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0", text):
        return text[:-2]
    return text or None


def _synthetic_id(*parts: Any) -> str:
    """Stable deduplication key derived from row content.

    Needed because the broker identifier is not always enough: on closed positions a
    single "Position ID" spans several rows when the holding was closed in parts
    (observed: 223 rows for 220 ids). Using the id alone as a unique key breaks the
    insert; ignoring it duplicates every re-import.
    """
    payload = "|".join("" if p is None else str(p) for p in parts)
    return "syn-" + hashlib.sha1(payload.encode("utf-8")).hexdigest()[:20]


def _build_open_positions(table: ParsedTable, result: ParsedExport) -> list[dict[str, Any]]:
    """Rebuild holdings from the aggregate rows.

    For each holding the export lists one aggregate row (category set, direction and
    open time blank) followed by one row per lot (direction and time set, category
    blank). Only aggregate rows become positions; summing both levels would count
    every holding twice.
    """
    aggregates: list[dict[str, Any]] = []
    lots_by_key: dict[tuple[str | None, str], list[dict[str, Any]]] = {}

    for row in table.rows:
        symbol = _clean(row.get("symbol"))
        if not symbol:
            continue
        symbol = symbol.upper()
        account = _clean(row.get("account"))

        is_lot = _clean(row.get("type")) is not None
        if is_lot:
            lots_by_key.setdefault((account, symbol), []).append(row)
        else:
            aggregates.append(row)

    # Fallback: some exports have no aggregate level. Each lot then becomes a position
    # in its own right rather than being lost.
    if not aggregates and lots_by_key:
        for lots in lots_by_key.values():
            aggregates.extend(lots)
        lots_by_key = {}

    positions: list[dict[str, Any]] = []

    for row in aggregates:
        symbol = str(_clean(row.get("symbol"))).upper()
        account = _clean(row.get("account"))
        volume = parse_number(row.get("volume"))
        open_price = parse_number(row.get("open_price"))

        if volume is None or open_price is None:
            result.warnings.append(Message(MessageCode.POSITION_SKIPPED, {"symbol": symbol}))
            continue

        lots = lots_by_key.get((account, symbol), [])
        lot_times = [parse_datetime(lot.get("open_time")) for lot in lots]
        lot_times = [t for t in lot_times if t is not None]

        # The current price only appears on lot rows. It is taken as-is — in the
        # instrument's currency — rather than derived from value / quantity, which
        # would yield a price in the account currency and break the comparison with
        # the average cost.
        market_price = parse_number(row.get("market_price"))
        if market_price is None:
            lot_prices = [parse_number(lot.get("market_price")) for lot in lots]
            market_price = next((p for p in lot_prices if p is not None), None)

        # The aggregate row carries no direction: infer it from the lots.
        directions = {_normalize(lot.get("type")) for lot in lots}
        if directions and directions <= {"sell", "vente", "short"}:
            volume = -abs(volume)

        positions.append(
            {
                "external_id": _clean_id(row.get("position_id"))
                or _synthetic_id(account, symbol, "open"),
                "broker_symbol": symbol,
                "name": _clean(row.get("name")),
                "category": _clean(row.get("category")),
                "account": account,
                "quantity": volume,
                "avg_price": open_price,
                "market_value": parse_number(row.get("market_value")),
                "market_price": market_price,
                "opened_at": min(lot_times) if lot_times else parse_datetime(row.get("open_time")),
                "lots_count": len(lots) or 1,
                "gross_pl": parse_number(row.get("gross_pl")),
                "net_pl": parse_number(row.get("net_pl")),
                "net_pl_pct": parse_number(row.get("net_pl_pct")),
                "purchase_value": parse_number(row.get("purchase_value")),
                "commission": parse_number(row.get("commission")),
                "swap": parse_number(row.get("swap")),
                "currency": _clean(row.get("currency")),
                "comment": _clean(row.get("comment")),
                "raw": row.get("_raw", {}),
            }
        )

    return positions


def _build_closed_positions(table: ParsedTable, result: ParsedExport) -> list[dict[str, Any]]:
    closed: list[dict[str, Any]] = []
    #: Occurrence counter per content key, so two byte-identical partial closes stay
    #: two distinct records.
    seen_keys: dict[str, int] = {}

    for row in table.rows:
        symbol = _clean(row.get("symbol"))
        if not symbol or is_total_row(row.get("name")) or is_total_row(symbol):
            continue
        symbol = symbol.upper()

        volume = parse_number(row.get("volume"))
        open_price = parse_number(row.get("open_price"))
        if volume is None or open_price is None:
            result.warnings.append(
                Message(MessageCode.CLOSED_POSITION_SKIPPED, {"symbol": symbol})
            )
            continue

        opened_at = parse_datetime(row.get("open_time"))
        closed_at = parse_datetime(row.get("close_time"))
        close_price = parse_number(row.get("close_price"))
        position_id = _clean_id(row.get("position_id"))

        # "Position ID" is not a sufficient key: a holding closed in several parts
        # produces several rows sharing one id. The key therefore combines the id
        # with the execution details.
        content_key = _synthetic_id(
            position_id, symbol, opened_at, closed_at, volume, open_price, close_price
        )
        occurrence = seen_keys.get(content_key, 0)
        seen_keys[content_key] = occurrence + 1
        external_id = content_key if occurrence == 0 else f"{content_key}-{occurrence}"

        closed.append(
            {
                "external_id": external_id,
                "position_id": position_id,
                "broker_symbol": symbol,
                "name": _clean(row.get("name")),
                "category": _clean(row.get("category")),
                "account": _clean(row.get("account")),
                "quantity": volume,
                "avg_price": open_price,
                "opened_at": opened_at,
                "closed_at": closed_at,
                "close_price": close_price,
                "purchase_value": parse_number(row.get("purchase_value")),
                "sale_value": parse_number(row.get("sale_value")),
                "open_fx_rate": parse_number(row.get("open_fx_rate")),
                "close_fx_rate": parse_number(row.get("close_fx_rate")),
                "gross_pl": parse_number(row.get("gross_pl")),
                "net_pl": parse_number(row.get("net_pl")),
                "commission": parse_number(row.get("commission")),
                "swap": parse_number(row.get("swap")),
                "currency": _clean(row.get("currency")),
                "comment": _clean(row.get("comment")),
                "raw": row.get("_raw", {}),
            }
        )

    return closed


def _build_cash_operations(table: ParsedTable) -> list[dict[str, Any]]:
    operations: list[dict[str, Any]] = []

    for row in table.rows:
        raw_type = row.get("type")
        # "Total" rows are subtotals, not operations.
        if is_total_row(raw_type):
            continue

        symbol = _clean(row.get("symbol"))
        executed_at = parse_datetime(row.get("time") or row.get("open_time"))
        amount = parse_number(row.get("amount"))

        operations.append(
            {
                "external_id": _clean_id(row.get("cash_id"))
                or _synthetic_id(raw_type, symbol, executed_at, amount),
                "type": classify_cash_type(raw_type),
                "raw_type": _clean(raw_type),
                "broker_symbol": symbol.upper() if symbol else None,
                "name": _clean(row.get("name")),
                "category": _clean(row.get("category")),
                "account": _clean(row.get("account")),
                "executed_at": executed_at,
                "amount": amount,
                "comment": _clean(row.get("comment")),
                "raw": row.get("_raw", {}),
            }
        )

    return operations
